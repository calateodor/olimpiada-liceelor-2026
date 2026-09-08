# -*- coding: utf-8 -*-
"""Verify PDFs: diacritics survive pdftotext, and every docx source line appears in the PDF text."""
import docx, json, os, re, subprocess, sys

BASE = "D:/Teo/PNL/Olimpiada/Olimpiada Liceelor 2026"
PDF_DIR = "D:/Teo/PNL/Site-uri/Olimpiada Liceelor/assets-src/regulamente/pdf"
J = json.load(open("D:/Teo/PNL/Site-uri/Olimpiada Liceelor/assets-src/regulamente/regulamente.json", encoding="utf-8"))
SRC = {r["slug"]: r["source"] for r in J["regulations"] if r["source_line_count"]}
CED = str.maketrans({"ş": "ș", "ţ": "ț", "Ş": "Ș", "Ţ": "Ț"})

def norm(s):
    s = s.translate(CED).replace("\u00a0", " ")
    return re.sub(r"\s+", " ", s).strip()

def pdf_text(path):
    out = subprocess.run(["pdftotext", "-enc", "UTF-8", "-layout", path, "-"], capture_output=True)
    return out.stdout.decode("utf-8", errors="replace")

def src_lines(slug):
    fname = SRC[slug]
    fname = "Regulamente/" + fname if not fname.startswith("anexa") else fname
    d = docx.Document(os.path.join(BASE, fname))
    lines = []
    for p in d.paragraphs:
        for ln in p.text.split("\n"):
            ln = norm(ln)
            if ln: lines.append(ln)
    for t in d.tables:
        for r in t.rows:
            for c in r.cells:
                for ln in c.text.split("\n"):
                    ln = norm(ln)
                    if ln: lines.append(ln)
    return lines

slugs = sys.argv[1:] or list(SRC)
total_missing = 0
for slug in slugs:
    txt = pdf_text(os.path.join(PDF_DIR, slug + ".pdf"))
    flat = norm(txt)
    bad = sorted(set(re.findall(r"[?\ufffd\u25a1]", txt)))
    dia = {ch: txt.count(ch) for ch in "șțăâîȘȚĂÂÎ"}
    lines = src_lines(slug)
    missing = []
    for ln in lines:
        # strip leading numbering / bullet chars that the layout may reformat
        core = re.sub(r"^(\d+(\.\d+)*\.?|[IVX]+\.|[a-z]\)|Art\.?\s*\d+(\.\d+)?\.?|\(\d+\)|Cros\s*[-–])\s*", "", ln)
        core = re.sub(r"[\s–-]+$", "", core)
        if core and core not in flat:
            # allow paragraph wrapping: check word-window of first 6 words
            w = core.split()
            if " ".join(w[:6]) not in flat:
                missing.append(ln[:90])
    total_missing += len(missing)
    print(f"{slug:15} pdf_chars={len(txt):6}  src_lines={len(lines):3}  missing={len(missing)}  "
          f"suspect_chars={bad}  diacritics={sum(dia.values())}")
    for m in missing: print("      MISSING:", m)
print("TOTAL MISSING:", total_missing)
